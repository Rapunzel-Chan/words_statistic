from typing import List
import openpyxl
from openpyxl.styles import Font, Alignment
import os
from ...domain.entities.words_statistic import WordStatistics


class ExcelGenerator:
    """Генератор Excel файлов"""

    def __init__(self, output_dir: str = "reports"):
        self.output_dir = output_dir
        os.makedirs(output_dir, exist_ok=True)

    async def generate(self, statistics: List[WordStatistics], filename: str) -> str:
        """Генерирует Excel файл со статистикой"""

        # Создаем новую книгу
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Word Statistics"

        # Заголовки
        headers = ["Словоформа", "Кол-во во всем документе", "Кол-во в каждой строке"]
        for col, header in enumerate(headers, 1):
            cell = ws.cell(row=1, column=col, value=header)
            cell.font = Font(bold=True)
            cell.alignment = Alignment(horizontal="center")

        # Данные
        for row, stat in enumerate(statistics, 2):
            ws.cell(row=row, column=1, value=stat.word_form)
            ws.cell(row=row, column=2, value=stat.total_count)

            # Формируем строку с количеством по строкам
            line_counts_str = ",".join(str(count) for count in stat.line_counts)
            ws.cell(row=row, column=3, value=line_counts_str)

            # Выравнивание
            ws.cell(row=row, column=3).alignment = Alignment(horizontal="left")

        # Автоподбор ширины колонок
        for column in ws.columns:
            max_length = 0
            column_letter = column[0].column_letter
            for cell in column:
                try:
                    if len(str(cell.value)) > max_length:
                        max_length = len(str(cell.value))
                except:
                    pass
            adjusted_width = min(max_length + 2, 100)
            ws.column_dimensions[column_letter].width = adjusted_width

        # Сохраняем файл
        output_path = os.path.join(self.output_dir, filename)
        wb.save(output_path)

        return output_path
